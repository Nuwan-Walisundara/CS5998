"""Stage 2 - Extraction.

Structure-aware extraction of requirement units from WSO2 integration requirement
documents that follow the organisational template (Markdown or .docx).

Output per document:
  units    - one row per template field or per requirement statement
  mappings - one row per source-to-target mapping row
  payload_keys - JSON keys found in the request payload (for consistency checks)
"""
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from . import config


@dataclass
class ExtractedDocument:
    doc: str
    units: list[dict] = field(default_factory=list)
    mappings: list[dict] = field(default_factory=list)
    payload_keys: set[str] = field(default_factory=set)


# --------------------------------------------------------------------------- helpers
def norm_key(s: str) -> str:
    """Normalise a template field / heading name for matching."""
    s = re.sub(r"[*\\:]", "", s or "")
    return re.sub(r"\s+", " ", s).strip().lower()


def clean_cell(s: str) -> str:
    s = (s or "").replace("\\_", "_").replace("\\<", "<").replace("\\>", ">").replace("\\.", ".")
    s = s.replace("**", "")
    return re.sub(r"\s+", " ", s).strip()


_NUMBERED = re.compile(r"(?:^|\s)\d{1,2}[.)]\s+(?=[A-Za-z])")
_SENTENCE = re.compile(r"(?<=[.!?])\s+(?=[A-Z])")


def split_statements(text: str) -> list[str]:
    """Split a cell into requirement statements: numbered list first, sentences as fallback."""
    if not text:
        return []
    if len(_NUMBERED.findall(text)) >= 2:
        parts = _NUMBERED.split(text)
    else:
        parts = _SENTENCE.split(text)
    return [p.strip(" ;") for p in parts if p and p.strip(" ;")]


_JSON_KEY = re.compile(r'"([A-Za-z_][\w-]*)"\s*:')


def payload_keys(text: str) -> set[str]:
    return set(_JSON_KEY.findall(text or ""))


# --------------------------------------------------------------------------- core
def _add_table(ed: ExtractedDocument, rows: list[list[str]], section: str, subsection: str):
    """Interpret one table (list of raw cell lists)."""
    rows = [r for r in rows if r and not all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in r)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return
    header = [norm_key(c) for c in rows[0]]

    # Mapping table: >= 4 columns with Source/Target headers
    if len(rows[0]) >= 4 and config.MAPPING_HEADER_KEYS <= {h.split()[0] for h in header if h}:
        cols = ["source", "source_field", "target", "target_field", "mandatory", "remarks"]
        for i, r in enumerate(rows[1:], 1):
            vals = [clean_cell(c) for c in r] + [""] * 6
            ed.mappings.append(dict(doc=ed.doc, section=section, mapping_table=subsection or section,
                                    row_no=i, **dict(zip(cols, vals[:6]))))
        return

    # Key/value table: first column is the field name
    for r in rows:
        if len(r) < 2:
            continue
        fname = clean_cell(r[0])
        value = clean_cell(" ".join(r[1:]))
        key = norm_key(fname)
        if not key:
            continue
        if key.startswith(config.PAYLOAD_FIELD_PREFIX):
            ed.payload_keys |= payload_keys(value)
        base = dict(doc=ed.doc, section=section, subsection=subsection, field=fname, field_key=key)
        if key in config.TEXT_REQUIREMENT_FIELDS and value:
            stmts = split_statements(value) if key in config.LIST_FIELDS else [value]
            for n, s in enumerate(stmts, 1):
                ed.units.append({**base, "item_no": n, "unit_type": "requirement", "text": s})
        else:
            ed.units.append({**base, "item_no": 0,
                             "unit_type": "requirement" if key in config.TEXT_REQUIREMENT_FIELDS else "field",
                             "text": value})


def parse_markdown(text: str, doc: str) -> ExtractedDocument:
    ed = ExtractedDocument(doc=doc)
    section = subsection = ""
    table: list[list[str]] = []

    def flush():
        nonlocal table
        if table:
            _add_table(ed, table, section, subsection)
        table = []

    for line in text.splitlines():
        s = line.strip()
        if s.startswith("|"):
            table.append([c for c in s.strip("|").split("|")])
            continue
        flush()
        if m := re.match(r"^(#{1,6})\s+(.*)", s):
            title = norm_key(m.group(2))
            if len(m.group(1)) <= 2:
                section, subsection = title, ""
            else:
                subsection = title
        elif re.fullmatch(r"\*\*[^*]+\*\*", s):  # bold line used as a heading
            subsection = norm_key(s)
    flush()
    return ed


def parse_docx(path: Path) -> ExtractedDocument:
    """Read headings and tables in body order from a Word document."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = Document(str(path))
    ed = ExtractedDocument(doc=path.stem)
    section = subsection = ""
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading"):
                level = int(re.sub(r"\D", "", style) or 1)
                if level <= 2:
                    section, subsection = norm_key(p.text), ""
                else:
                    subsection = norm_key(p.text)
        elif tag == "tbl":
            t = Table(child, d)
            rows = [[c.text for c in row.cells] for row in t.rows]
            _add_table(ed, rows, section, subsection)
    return ed


def extract_file(path: Path) -> ExtractedDocument:
    path = Path(path)
    if path.suffix.lower() == ".docx":
        return parse_docx(path)
    return parse_markdown(path.read_text(encoding="utf-8"), path.stem)


def to_frames(docs: list[ExtractedDocument]) -> tuple[pd.DataFrame, pd.DataFrame]:
    units = pd.DataFrame([u for d in docs for u in d.units])
    mappings = pd.DataFrame([m for d in docs for m in d.mappings])
    if not units.empty:
        units.insert(0, "req_id", [f"R{i:04d}" for i in range(1, len(units) + 1)])
    return units, mappings


def main():
    ap = argparse.ArgumentParser(description="Extract requirement units from documents")
    ap.add_argument("--input", default=str(config.RAW_DIR))
    ap.add_argument("--out", default=str(config.PROCESSED_DIR))
    a = ap.parse_args()
    files = [p for p in Path(a.input).glob("*") if p.suffix.lower() in {".md", ".docx"}]
    docs = [extract_file(p) for p in sorted(files)]
    units, mappings = to_frames(docs)
    out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    units.to_csv(out / "requirements_raw.csv", index=False)
    mappings.to_csv(out / "mappings_raw.csv", index=False)
    print(f"{len(docs)} documents -> {len(units)} units, {len(mappings)} mapping rows -> {out}")


if __name__ == "__main__":
    main()
